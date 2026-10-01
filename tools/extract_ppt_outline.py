# -*- coding: utf-8 -*-
"""Extract full text + structure + images from the competition PPTX."""
import os, sys, json, io
from pptx import Presentation
from pptx.util import Emu

SRC = r"C:\Users\ac010\Desktop\医数智答——医疗大数据知识问答系统.pptx"
OUT = r"C:\Users\ac010\Documents\deepseek-harness\default-workspace\_ppt_dump"
os.makedirs(OUT, exist_ok=True)
IMGDIR = os.path.join(OUT, "images")
os.makedirs(IMGDIR, exist_ok=True)

prs = Presentation(SRC)
print("SLIDE SIZE:", prs.slide_width, prs.slide_height, Emu(prs.slide_width).inches, Emu(prs.slide_height).inches)
print("SLIDES:", len(prs.slides))
print("=" * 100)

lines = []
img_index = 0

def walk(shapes, depth=0, slide_no=0):
    global img_index
    out = []
    for sh in shapes:
        pad = "  " * depth
        t = sh.shape_type
        name = sh.name
        info = f"{pad}[{t}] name={name}"
        try:
            info += f" pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) size=({Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})"
        except Exception:
            pass
        out.append(info)
        if sh.shape_type == 6 or sh.__class__.__name__ == "GroupShape":
            try:
                out.extend(walk(sh.shapes, depth + 1, slide_no))
                continue
            except Exception:
                pass
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                txt = "".join(r.text for r in p.runs)
                if txt.strip():
                    out.append(f"{pad}  TXT(lvl{p.level}): {txt}")
        if sh.has_table:
            tbl = sh.table
            for r_i, row in enumerate(tbl.rows):
                cells = [c.text.replace("\n", " / ") for c in row.cells]
                out.append(f"{pad}  ROW{r_i}: " + " | ".join(cells))
        if sh.shape_type == 13 or sh.__class__.__name__ == "Picture":
            try:
                img = sh.image
                ext = img.ext
                img_index += 1
                fn = os.path.join(IMGDIR, f"s{slide_no:02d}_{img_index:03d}.{ext}")
                with open(fn, "wb") as f:
                    f.write(img.blob)
                out.append(f"{pad}  IMAGE -> {fn} ({len(img.blob)} bytes)")
            except Exception as e:
                out.append(f"{pad}  IMAGE err {e}")
        if getattr(sh, "has_chart", False) and sh.has_chart:
            try:
                ch = sh.chart
                out.append(f"{pad}  CHART type={ch.chart_type}")
                for s in ch.plots[0].series:
                    out.append(f"{pad}    series {s.name}: {list(s.values)}")
                out.append(f"{pad}    cats {list(ch.plots[0].categories)}")
            except Exception as e:
                out.append(f"{pad}  CHART err {e}")
    return out

for i, slide in enumerate(prs.slides, 1):
    header = f"\n{'='*40} SLIDE {i} {'='*40}"
    print(header)
    lines.append(header)
    body = walk(slide.shapes, 0, i)
    for b in body:
        print(b)
        lines.append(b)
    if slide.has_notes_slide:
        nt = slide.notes_slide.notes_text_frame.text
        if nt.strip():
            print("--- NOTES:", nt)
            lines.append("--- NOTES: " + nt)

with open(os.path.join(OUT, "dump.txt"), "w", encoding="utf-8") as f:
    f.write("\n".join(lines))
print("\nSAVED to", os.path.join(OUT, "dump.txt"))
